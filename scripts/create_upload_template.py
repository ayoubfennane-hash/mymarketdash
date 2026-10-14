#!/usr/bin/env python3
"""Crée un modèle Excel vierge conforme à la feuille de données importée par MYMARKET."""
from __future__ import annotations
import argparse
from copy import copy
from pathlib import Path
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT=Path(__file__).resolve().parents[1]
DEFAULT_SOURCE=ROOT/'public'/'MYMARKET_Tableau_de_bord.xlsx'
DEFAULT_OUTPUT=ROOT/'public'/'MYMARKET_modele_import.xlsx'
GREEN='12645C'; DARK='203D33'; PALE='F0F5EF'; LIGHT='F7F9F5'; GOLD='E6AD46'; WHITE='FFFFFF'; BLUE='1768AC'; LINE='DCE6DD'

def main(source:Path, output:Path):
    src=openpyxl.load_workbook(source,read_only=True,data_only=True)
    data_sheet=None; headers=None
    for sheet in src.worksheets:
        row=next(sheet.iter_rows(min_row=1,max_row=1,values_only=True),())
        vals=list(row[:28])
        if len(vals)==28 and str(vals[0]).strip().casefold()=='title' and str(vals[1]).strip().upper()=='SKU' and str(vals[2]).strip().casefold()=='fournisseur' and str(vals[5]).strip().casefold()=='situation':
            data_sheet=sheet.title;headers=vals;break
    if headers is None: raise ValueError('Feuille Excel source introuvable : 28 colonnes attendues, Title/SKU/fournisseur/situation en tête.')
    if any(not hasattr(headers[i],'year') for i in list(range(6,16))+list(range(17,27))): raise ValueError('Les deux séries de mois doivent être des dates Excel.')
    if [headers[i].date() for i in range(6,16)] != [headers[i].date() for i in range(17,27)]: raise ValueError('Les mois des quantités et montants ne correspondent pas.')
    src.close()
    wb=Workbook(); ws=wb.active; ws.title='Feuil1'; info=wb.create_sheet('Instructions')
    ws.sheet_view.showGridLines=False; ws.freeze_panes='G2'; ws.sheet_properties.tabColor=GREEN
    ws.append(headers)
    header_fill=PatternFill('solid',fgColor=GREEN); total_fill=PatternFill('solid',fgColor='E8F0E9')
    input_font=Font(name='Aptos',size=10,color=BLUE); formula_font=Font(name='Aptos',size=10,color=DARK)
    for i,value in enumerate(headers,1):
        c=ws.cell(1,i);c.value=value;c.font=Font(name='Aptos',size=10,bold=True,color=WHITE);c.fill=header_fill;c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        if 7<=i<=16 or 18<=i<=27:c.number_format='mmm yyyy'
    widths={1:42,2:22,3:26,4:28,5:23,6:21,17:15,28:15}
    for col,width in widths.items():ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width=width
    for col in list(range(7,17))+list(range(18,28)):ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width=14
    ws.row_dimensions[1].height=38
    for r in range(2,102):
        for col in range(1,29):
            cell=ws.cell(r,col);cell.fill=PatternFill('solid',fgColor=WHITE if r%2==0 else LIGHT);cell.border=Border(bottom=Side(style='hair',color=LINE));cell.alignment=Alignment(vertical='center')
            if col not in (17,28):cell.font=input_font
        ws.cell(r,2).number_format='@'
        ws.cell(r,17,f'=SUM(G{r}:P{r})');ws.cell(r,28,f'=SUM(R{r}:AA{r})')
        for col in (17,28):
            ws.cell(r,col).font=formula_font;ws.cell(r,col).fill=total_fill
            ws.cell(r,col).number_format='#,##0.##;[Red]-#,##0.##;-' if col==17 else '#,##0.00;[Red]-#,##0.00;-'
        ws.cell(r,5).number_format='#,##0.##;[Red]-#,##0.##;-'
        for col in list(range(7,17)):ws.cell(r,col).number_format='#,##0.##;[Red]-#,##0.##;-'
        for col in list(range(18,28)):ws.cell(r,col).number_format='#,##0.00;[Red]-#,##0.00;-'
        ws.row_dimensions[r].height=20
    ws.auto_filter.ref='A1:AB101'
    numeric=DataValidation(type='decimal',operator='between',formula1='-1000000000',formula2='1000000000',allow_blank=True);numeric.errorTitle='Valeur non numérique';numeric.error='Saisissez un nombre. Les valeurs négatives sont admises.';numeric.showErrorMessage=True;ws.add_data_validation(numeric)
    for col in ['E','G','H','I','J','K','L','M','N','O','P','R','S','T','U','V','W','X','Y','Z','AA']:numeric.add(f'{col}2:{col}101')
    status=DataValidation(type='list',formula1='"En stock,Presque rupture,Rupture,Non renseigné"',allow_blank=True);status.error='Choisissez un état proposé ou collez une autre valeur.';status.showErrorMessage=True;ws.add_data_validation(status);status.add('F2:F101')
    ws.print_title_rows='1:1';ws.print_area='A1:AB101';ws.page_setup.orientation='landscape';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.pageSetUpPr.fitToPage=True
    info.sheet_view.showGridLines=False;info.sheet_properties.tabColor=GOLD;info.column_dimensions['A'].width=24;info.column_dimensions['B'].width=70
    for col in 'CDEF':info.column_dimensions[col].width=6
    info.merge_cells('A1:F1');info['A1']='MYMARKET — Modèle Excel à compléter';info['A1'].font=Font(name='Aptos Display',size=18,bold=True,color=WHITE);info['A1'].fill=PatternFill('solid',fgColor=GREEN);info['A1'].alignment=Alignment(vertical='center');info.row_dimensions[1].height=38
    lines=[
      ('Objectif','Saisissez vos références dans la feuille Feuil1, une ligne par produit ou variante. Les 100 lignes sont vierges; aucune donnée produit fictive n’est fournie.'),
      ('Colonnes A:F','Title, SKU, fournisseur, Location, Available (not editable), situation. Ne renommez pas et ne déplacez pas ces en-têtes.'),
      ('SKU','Conservez le SKU comme texte, sans supprimer les zéros initiaux.'),
      ('Quantités G:P','Saisissez les quantités mensuelles. Les dix mois indiqués dans l’en-tête sont des dates Excel.'),
      ('Total quantités Q','Calculé automatiquement par formule à partir des colonnes G:P. Ne remplacez pas la formule.'),
      ('Montants R:AA','Saisissez les montants mensuels correspondants aux mêmes mois que les quantités.'),
      ('Total montants AB','Calculé automatiquement par formule à partir des colonnes R:AA. Ne remplacez pas la formule.'),
      ('Mois','Si vous changez les mois, modifiez les en-têtes des deux séries. Elles doivent contenir exactement les mêmes dix dates.'),
      ('Type de produit','La visualisation rapproche les catégories depuis le CSV fourni, par SKU exact puis titre exact. Il n’est pas nécessaire d’ajouter une colonne catégorie au modèle.'),
      ('Import','Conservez les 28 colonnes dans le même ordre, remplissez Title pour chaque référence et gardez les cellules numériques sous forme de nombres.'),
    ]
    for row,(label,text) in enumerate(lines,3):
        a=info.cell(row,1,label);a.font=Font(name='Aptos',size=10,bold=True,color=DARK);a.fill=PatternFill('solid',fgColor=PALE);a.alignment=Alignment(vertical='top',wrap_text=True)
        b=info.cell(row,2,text);b.font=Font(name='Aptos',size=10,color=DARK);b.alignment=Alignment(vertical='top',wrap_text=True);info.merge_cells(start_row=row,start_column=2,end_row=row,end_column=6);info.row_dimensions[row].height=48
    info.freeze_panes='A3';info.sheet_properties.pageSetUpPr.fitToPage=True;info.page_setup.orientation='landscape';info.page_setup.fitToWidth=1;info.page_setup.fitToHeight=1;info.sheet_properties.tabColor=GOLD
    wb.active=0
    if hasattr(wb,'calculation'):
        wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True;wb.calculation.calcMode='auto'
    output.parent.mkdir(parents=True,exist_ok=True);wb.save(output)
    print(f'Modèle créé : {output} · feuille source {data_sheet} · 28 colonnes · 100 lignes vides avec formules Q/AB.')

if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--source',type=Path,default=DEFAULT_SOURCE);p.add_argument('--output',type=Path,default=DEFAULT_OUTPUT);a=p.parse_args();main(a.source,a.output)
